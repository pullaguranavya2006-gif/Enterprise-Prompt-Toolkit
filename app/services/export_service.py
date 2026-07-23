import json

from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from docx import Document


def export_pdf(prompt, filename):

    pdf = canvas.Canvas(filename, pagesize=letter)

    pdf.setFont("Helvetica", 14)

    pdf.drawString(40, 760, "Enterprise Prompt Engineering Toolkit")

    pdf.setFont("Helvetica", 11)

    y = 730

    for line in prompt.split("\n"):

        pdf.drawString(40, y, line)

        y -= 20

        if y < 50:
            pdf.showPage()
            y = 760

    pdf.save()


def export_docx(prompt, filename):

    document = Document()

    document.add_heading(
        "Enterprise Prompt Engineering Toolkit",
        level=1,
    )

    document.add_paragraph(prompt)

    document.save(filename)


def export_json(data, filename):

    with open(filename, "w") as file:

        json.dump(
            data,
            file,
            indent=4
        )